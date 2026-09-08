"""Regenerate the synthetic XLSX companion from the checked-in CSV."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

root = Path(__file__).resolve().parents[1]
source = root / 'data/samples/project-snapshot.csv'
book = Workbook()
sheet = book.active
sheet.title = 'Project Snapshot'
with source.open(encoding='utf-8', newline='') as handle:
    for row in csv.reader(handle):
        sheet.append(row)
for cell in sheet[1]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='173E35')
sheet.freeze_panes = 'A2'
sheet.auto_filter.ref = sheet.dimensions
for column, width in {'A':18,'B':30,'C':20,'D':20,'E':20,'F':20,'G':20,'H':20}.items():
    sheet.column_dimensions[column].width = width
book.save(root / 'data/samples/project-snapshot.xlsx')
print('Synthetic XLSX sample generated.')

