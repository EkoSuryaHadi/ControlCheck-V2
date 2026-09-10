import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from controlcheck.main import create_app


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path / 'test.db')) as c:
        yield c


def project(client, name='Test'):
    response = client.post('/api/projects', json={'name':name,'currency':'IDR','as_of':'2026-09-08'})
    assert response.status_code == 201
    return response.json()['id']


def upload(client, pid, content=None):
    content = content or b'Activity ID,Name,Planned Progress,Actual Progress,Budget,Actual Cost\nA1,Foundation,80,40,1000,600'
    response = client.post(f'/api/projects/{pid}/sources', files={'file':('sample.csv', content,'text/csv')})
    assert response.status_code == 201, response.text
    return response.json()


def test_full_journey_and_idempotency(client):
    pid = project(client); source = upload(client,pid)
    mapping = {s['column']:s['field'] for s in source['suggestions'] if s['field']}
    root = f"/api/projects/{pid}/sources/{source['id']}"
    preview = client.post(root+'/validate',json={'mapping':mapping})
    assert preview.status_code == 200 and preview.json()['errors'] == []
    response = client.post(root+'/publish',json={'mapping':mapping})
    assert response.status_code == 200
    assert client.post(root+'/publish',json={'mapping':mapping}).json()['id'] == response.json()['id']
    overview = client.get(f'/api/projects/{pid}/overview').json()
    assert overview['analysis']['metrics']['spi'] == .5
    answer = client.post(f'/api/projects/{pid}/assistant',json={'question':'Bagaimana kondisi proyek?'})
    assert answer.json()['evidence'][0]['row'] == 2
    assert 'sample.csv' in client.get(f'/api/projects/{pid}/report').text


def test_project_isolation_and_no_data(client):
    a,b = project(client,'A'),project(client,'B')
    source = upload(client,b)
    assert client.post(f"/api/projects/{a}/sources/{source['id']}/publish", json={'mapping':{}}).status_code == 404
    assert client.get(f'/api/projects/{a}/sources').json() == []
    assert client.post(f'/api/projects/{a}/assistant',json={'question':'Kondisi?'}).status_code == 409


def test_invalid_publish_does_not_activate(client):
    pid = project(client)
    source = upload(client,pid,b'Activity ID,Name,Actual Progress\nA1,Test,120')
    mapping = {s['column']:s['field'] for s in source['suggestions']}
    assert client.post(f"/api/projects/{pid}/sources/{source['id']}/publish",json={'mapping':mapping}).status_code == 422
    assert client.get(f'/api/projects/{pid}/overview').json()['snapshot'] is None


def test_xlsx_selected_sheet_and_formula_rejection(client):
    pid = project(client)
    book = Workbook(); book.active.title='Cover'; ws=book.create_sheet('Schedule')
    ws.append(['Activity ID','Name']); ws.append(['A1','Work'])
    stream=io.BytesIO(); book.save(stream)
    route=f'/api/projects/{pid}/sources?sheet=Schedule'
    result=client.post(route,files={'file':('test.xlsx',stream.getvalue())})
    assert result.status_code == 201 and result.json()['sheet'] == 'Schedule'
    ws['B2']='=1+1'; stream=io.BytesIO(); book.save(stream)
    assert client.post(route,files={'file':('test.xlsx',stream.getvalue())}).status_code == 422


def test_reporting_date_and_unsupported_type(client):
    assert client.post('/api/projects',json={'name':'X','currency':'IDR','as_of':'yesterday'}).status_code == 422
    pid=project(client)
    assert client.post(f'/api/projects/{pid}/sources',files={'file':('a.xer',b'anything')}).status_code == 422


def test_replacement_and_unsupported_answer(client):
    pid = project(client)
    for expected, data in [(1,b'Activity ID,Name\nA1,First'),(2,b'Activity ID,Name\nB1,Second')]:
        s=upload(client,pid,data)
        mapping={r['column']:r['field'] for r in s['suggestions']}
        result=client.post(f"/api/projects/{pid}/sources/{s['id']}/publish",json={'mapping':mapping})
        assert result.json()['version'] == expected
    overview=client.get(f'/api/projects/{pid}/overview').json()
    assert [r['activity_id'] for r in overview['snapshot']['rows']] == ['B1']
    answer=client.post(f'/api/projects/{pid}/assistant',json={'question':'Kenapa SPI turun?'}).json()
    assert answer['evidence'] == [] and 'belum cukup' in answer['answer']


def test_persistence_after_restart(tmp_path):
    db=tmp_path/'persistent.db'
    with TestClient(create_app(db)) as client:
        pid=project(client)
        upload(client,pid)
    with TestClient(create_app(db)) as client:
        assert client.get('/api/projects').json()[0]['id'] == pid
        assert len(client.get(f'/api/projects/{pid}/sources').json()) == 1


def test_inspect_then_upload_with_header_and_normalization(client):
    pid=project(client)
    content=b'Report title;\nActivity ID;Name;Planned Finish;Actual Progress;Budget\nA1;Work;31/12/2026;0,4;1000,5'
    inspected=client.post(f'/api/projects/{pid}/sources/inspect',files={'file':('report.csv',content)})
    assert inspected.status_code == 200
    assert inspected.json()['sheets'][0]['suggested_header_row'] == 2
    route=(f'/api/projects/{pid}/sources?header_row=2&date_format=dmy'
           '&decimal_separator=comma&percent_scale=fraction')
    uploaded=client.post(route,files={'file':('report.csv',content)})
    assert uploaded.status_code == 201, uploaded.text
    source=uploaded.json(); mapping={s['column']:s['field'] for s in source['suggestions'] if s['field']}
    preview=client.post(f"/api/projects/{pid}/sources/{source['id']}/validate",json={'mapping':mapping})
    assert preview.status_code == 200 and preview.json()['errors'] == []


def test_progress_source_uploads_with_dataset_type(client):
    pid = project(client)
    response = client.post(
        f'/api/projects/{pid}/sources?dataset_type=progress',
        files={'file': ('progress.csv', b'Activity ID,Actual Progress\nA1,55', 'text/csv')},
    )
    assert response.status_code == 201, response.text
    source = response.json()
    assert source['dataset_type'] == 'progress'
    mapping = {item['column']: item['field'] for item in source['suggestions'] if item['field']}
    validated = client.post(f"/api/projects/{pid}/sources/{source['id']}/validate", json={'mapping': mapping})
    assert validated.status_code == 200
    assert validated.json()['errors'] == []


def test_schedule_reconciliation_preview_accepts_schedule_only(client):
    pid = project(client)
    uploaded = client.post(
        f'/api/projects/{pid}/sources?dataset_type=schedule',
        files={'file': ('schedule.csv', b'Activity ID,Name\nA1,Foundation', 'text/csv')},
    )
    source = uploaded.json()
    mapping = {item['column']: item['field'] for item in source['suggestions'] if item['field']}
    preview = client.post(f'/api/projects/{pid}/reconciliations', json={
        'schedule_source_id': source['id'], 'schedule_mapping': mapping,
    })
    assert preview.status_code == 200, preview.text
    assert preview.json()['coverage'] == {'schedule': 1, 'progress_matched': 0, 'cost_matched': 0}


def test_schedule_reconciliation_publish_activates_schedule_only_snapshot(client):
    pid = project(client)
    uploaded = client.post(
        f'/api/projects/{pid}/sources?dataset_type=schedule',
        files={'file': ('schedule.csv', b'Activity ID,Name\nA1,Foundation', 'text/csv')},
    )
    source = uploaded.json()
    mapping = {item['column']: item['field'] for item in source['suggestions'] if item['field']}
    payload = {'schedule_source_id': source['id'], 'schedule_mapping': mapping}
    published = client.post(f'/api/projects/{pid}/reconciliations/publish', json=payload)
    assert published.status_code == 200, published.text
    assert published.json()['source_ids'] == [source['id']]
    assert published.json()['semantic_version'] == 'activity-snapshot/v2'
    overview = client.get(f'/api/projects/{pid}/overview').json()
    assert overview['snapshot']['id'] == published.json()['id']


def test_agent_upload_publishes_valid_schedule_without_mapping(client):
    pid = project(client)
    response = client.post(f'/api/projects/{pid}/ingestions', files=[
        ('files', ('schedule.csv', b'Activity ID,Name\nA1,Foundation', 'text/csv')),
    ])
    assert response.status_code == 201, response.text
    assert response.json()['status'] == 'published'
    assert response.json()['snapshot']['rows'][0]['activity_id'] == 'A1'


def test_agent_does_not_publish_unmatched_progress_activity(client):
    pid = project(client)
    response = client.post(f'/api/projects/{pid}/ingestions', files=[
        ('files', ('schedule.csv', b'Activity ID,Name\nA1,Foundation', 'text/csv')),
        ('files', ('progress.csv', b'Activity ID,Actual Progress\nB9,55', 'text/csv')),
    ])
    assert response.status_code == 200, response.text
    assert response.json()['status'] == 'needs_attention'
    assert client.get(f'/api/projects/{pid}/overview').json()['snapshot'] is None


def test_assistant_reports_dependency_impact_when_schedule_supplies_links(client):
    pid = project(client)
    content = (b'Activity ID,Name,Planned Finish,Actual Progress,Is Critical,Is Milestone,Predecessor IDs\n'
               b'A,Late critical,2026-09-01,20,true,false,\n'
               b'B,Downstream milestone,2026-09-30,0,false,true,A')
    source = upload(client, pid, content)
    mapping = {item['column']: item['field'] for item in source['suggestions'] if item['field']}
    assert client.post(f"/api/projects/{pid}/sources/{source['id']}/publish", json={'mapping': mapping}).status_code == 200
    answer = client.post(f'/api/projects/{pid}/assistant', json={'question': 'Apa dampak dependency critical?'}).json()
    assert 'jalur dampak: 1' in answer['answer']
    assert len(answer['evidence']) == 2
